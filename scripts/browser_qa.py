"""Local QA against a disposable database. Requires Playwright + Chromium and openpyxl."""

import csv
import json
import os
import tempfile
from io import StringIO
from pathlib import Path

from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright

ROOT = Path(__file__).resolve().parents[1]
BASE = os.environ.get("QA_BASE_URL", "http://127.0.0.1:5173")
SHOTS = ROOT / "docs/screenshots"


def run():
    SHOTS.mkdir(parents=True, exist_ok=True)
    results = {"viewports": [], "accessibility": [], "downloads": [], "errors": []}
    with sync_playwright() as playwright, tempfile.TemporaryDirectory() as temporary:
        browser = playwright.chromium.launch()
        page = browser.new_page(viewport={"width": 1440, "height": 1000})
        page.on("pageerror", lambda error: results["errors"].append(str(error)))

        def navigate(route, heading):
            page.goto(BASE + "/#" + route)
            expect(page.get_by_role("heading", name=heading, exact=True)).to_be_visible()
            page.wait_for_timeout(400)

        def download(name):
            with page.expect_download() as event:
                page.get_by_role("button", name=name, exact=True).click()
            saved = Path(temporary) / event.value.suggested_filename
            event.value.save_as(saved)
            results["downloads"].append(saved.name)
            return saved

        initial = page.request.get(BASE + "/api/imports").json()["total"]
        assert initial == 0, (
            "Use a fresh disposable QA database; this script creates fictional imports."
        )
        navigate("/", "Good data. Clear decisions.")
        page.screenshot(path=str(SHOTS / "01-overview.png"), full_page=True)
        page.get_by_role("link", name="Explore demo", exact=True).first.click()
        expect(page.get_by_role("heading", name="Processing report")).to_be_visible()
        expect(page.get_by_text("3,445.00", exact=True)).to_be_visible()
        page.wait_for_timeout(600)
        page.screenshot(path=str(SHOTS / "02-sales-report.png"), full_page=True)
        navigate("/demo/issues", "Your sales overview")
        page.get_by_role("tab", name="Issues").click()
        expect(page.get_by_text("normalized_duplicate", exact=True)).to_be_visible()
        page.locator(".report-card").screenshot(path=str(SHOTS / "03-issues.png"))
        issues = list(csv.DictReader(StringIO(download("Issues CSV").read_text("utf-8-sig"))))
        assert len(issues) == 8
        assert set(issues[0]) == {"source_row", "column", "code", "explanation"}
        page.get_by_role("tab", name="Issues").focus()
        page.keyboard.press("ArrowRight")
        expect(page.get_by_role("tab", name="Changes")).to_be_focused()
        expect(page.get_by_text("Trimmed surrounding whitespace.")).to_be_visible()
        page.locator(".report-card").screenshot(path=str(SHOTS / "04-changes.png"))
        assert page.request.get(BASE + "/api/imports").json()["total"] == initial
        source = page.request.get(BASE + "/api/demos/issues/source").body()

        navigate("/import", "Import your sales data")
        page.get_by_label("File currency").select_option("USD")
        page.locator('input[type="file"]').set_input_files(
            {
                "name": "fictional-stationery-issues.csv",
                "mimeType": "text/csv",
                "buffer": source,
            }
        )
        page.screenshot(path=str(SHOTS / "05-import.png"), full_page=True)
        page.get_by_role("button", name="Import & process").click()
        expect(page.get_by_role("heading", name="Processing report")).to_be_visible()
        report_route = page.url.split("#")[1]
        clean_csv = download("CSV")
        assert len(list(csv.DictReader(StringIO(clean_csv.read_text("utf-8-sig"))))) == 48
        clean_xlsx = download("XLSX")
        workbook = load_workbook(clean_xlsx, data_only=False)
        assert workbook.active.max_row == 49
        workbook.close()
        assert "USD" in clean_csv.name and "USD" in clean_xlsx.name
        assert (
            len(list(csv.DictReader(StringIO(download("Issues CSV").read_text("utf-8-sig"))))) == 8
        )

        # A clean Excel export can pass through the upload workflow with an explicit new currency.
        navigate("/import", "Import your sales data")
        page.locator('input[type="file"]').set_input_files(str(clean_xlsx))
        expect(page.get_by_role("button", name="Import & process")).to_be_disabled()
        page.get_by_label("File currency").select_option("EUR")
        page.get_by_role("button", name="Import & process").click()
        expect(page.get_by_role("heading", name="Processing report")).to_be_visible()
        expect(page.get_by_text("Total revenue · EUR", exact=True)).to_be_visible()
        assert (
            page.request.get(BASE + "/api" + page.url.split("#")[1] + "/analytics").json()[
                "currency"
            ]
            == "EUR"
        )
        assert "EUR" in download("CSV").name

        navigate("/import", "Import your sales data")
        page.get_by_label("File currency").select_option("PLN")
        page.locator('input[type="file"]').set_input_files(
            {
                "name": "fictional-invalid.csv",
                "mimeType": "text/csv",
                "buffer": (
                    b"order_id,date,customer,product,quantity,unit_price\n1,bad,Mira,Book,0,-1"
                ),
            }
        )
        page.get_by_role("button", name="Import & process").click()
        expect(page.get_by_role("link", name="Fix file & import again")).to_be_visible()
        expect(page.get_by_role("button", name="CSV", exact=True)).to_be_disabled()
        assert (
            len(list(csv.DictReader(StringIO(download("Issues CSV").read_text("utf-8-sig"))))) == 3
        )
        page.get_by_role("link", name="Fix file & import again").click()
        expect(page.get_by_role("heading", name="Import your sales data")).to_be_visible()
        page.get_by_label("File currency").select_option("USD")
        page.locator('input[type="file"]').set_input_files(
            {
                "name": "bad-columns.csv",
                "mimeType": "text/csv",
                "buffer": b"wrong,columns\n1,2",
            }
        )
        page.get_by_role("button", name="Import & process").click()
        expect(page.get_by_role("alert")).to_contain_text("columns")

        # History failure must not prevent opening a report or exporting its rows.
        def fail_history(route):
            route.fulfill(
                status=503,
                content_type="application/json",
                body=json.dumps(
                    {
                        "detail": {"message": "Simulated history outage."},
                    }
                ),
            )

        page.route("**/api/imports?*", fail_history)
        navigate(report_route, "Your sales overview")
        expect(page.get_by_role("heading", name="Processing report")).to_be_visible()
        expect(page.get_by_role("alert")).to_contain_text("History could not be loaded")
        download("CSV")
        page.unroute("**/api/imports?*", fail_history)

        # Seed only the disposable QA database, enough to exercise history pagination.
        for index in range(9):
            response = page.request.post(
                BASE + "/api/imports",
                multipart={
                    "currency": "USD",
                    "file": {
                        "name": f"fictional-history-check-{index + 1:02}.csv",
                        "mimeType": "text/csv",
                        "buffer": source,
                    },
                },
            )
            assert response.status == 201
        navigate("/history", "Processing history")
        page.get_by_role("button", name="Next imports").click()
        expect(page.get_by_role("button", name="Previous imports")).to_be_enabled()
        page.get_by_role("link", name="Open fictional-stationery-issues.csv", exact=True).click()
        expect(page.get_by_role("heading", name="Processing report")).to_be_visible()

        axe = os.environ.get("AXE_PATH")
        for width in [320, 375, 768, 1440]:
            page.set_viewport_size({"width": width, "height": 1000})
            for route, heading in [
                ("/", "Good data. Clear decisions."),
                ("/demo/issues", "Your sales overview"),
                ("/import", "Import your sales data"),
                ("/history", "Processing history"),
            ]:
                navigate(route, heading)
                assert page.evaluate("document.documentElement.scrollWidth <= innerWidth"), (
                    width,
                    route,
                )
                results["viewports"].append({"width": width, "route": route, "overflow": False})
                if route == "/demo/issues":
                    for tab in ["Issues", "Changes", "Cleaned data"]:
                        page.get_by_role("tab", name=tab).click()
                        assert page.evaluate(
                            "document.documentElement.scrollWidth <= innerWidth"
                        ), (width, tab)
                if axe:
                    page.add_script_tag(path=axe)
                    violations = page.evaluate(
                        "async () => (await axe.run()).violations.map(v => "
                        "({id:v.id, impact:v.impact, nodes:v.nodes.map(n=>n.target)}))"
                    )
                    results["accessibility"].append(
                        {"width": width, "route": route, "violations": violations}
                    )

        page.set_viewport_size({"width": 375, "height": 900})
        navigate("/demo/issues", "Your sales overview")
        page.screenshot(path=str(SHOTS / "06-mobile-report.png"), full_page=True)
        page.get_by_role("button", name="Open navigation").click()
        expect(page.get_by_role("link", name="FlowOps home")).to_be_focused()
        page.keyboard.press("Escape")
        expect(page.get_by_role("button", name="Open navigation")).to_be_focused()
        browser.close()
    report = ROOT / "artifacts/browser-qa.json"
    report.parent.mkdir(exist_ok=True)
    report.write_text(json.dumps(results, indent=2), encoding="utf-8")
    print(json.dumps(results, indent=2))
    assert not results["errors"]
    assert not any(item["violations"] for item in results["accessibility"])


if __name__ == "__main__":
    run()
