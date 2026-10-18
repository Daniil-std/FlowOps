"""Exercise the built browser app in an isolated Chromium context; no backend needed."""

import csv
import json
import os
import re
import tempfile
import time
from io import StringIO
from pathlib import Path
from urllib.request import urlopen

from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright

ROOT = Path(__file__).resolve().parents[1]
BASE = os.environ.get("QA_BASE_URL", "http://127.0.0.1:5174").rstrip("/")
OUTPUT = ROOT / "artifacts/browser-demo"


def run():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    for attempt in range(30):
        try:
            with urlopen(BASE, timeout=2) as response:
                assert response.status == 200
            break
        except OSError:
            if attempt == 29:
                raise
            time.sleep(0.5)
    results = {
        "errors": [],
        "workers": [],
        "viewports": [],
        "accessibility": [],
        "checks": [],
    }
    with sync_playwright() as playwright, tempfile.TemporaryDirectory() as temporary:
        browser = playwright.chromium.launch()
        context = browser.new_context(viewport={"width": 1440, "height": 1000})
        page = context.new_page()
        page.on("pageerror", lambda error: results["errors"].append(str(error)))
        page.on("worker", lambda worker: results["workers"].append(worker.url))
        requests = []
        page.on("request", lambda request: requests.append((request.method, request.url)))

        def navigate(route, heading):
            page.goto(BASE + "/#" + route)
            expect(page.get_by_role("heading", name=heading, exact=True)).to_be_visible()

        def select_file(content, name="my-orders.csv", currency="USD"):
            navigate("/import", "Import your sales data")
            page.get_by_label("File currency", exact=True).select_option(currency)
            page.locator('input[type="file"]').set_input_files(
                {
                    "name": name,
                    "mimeType": "application/octet-stream",
                    "buffer": content,
                }
            )

        def process():
            page.get_by_role("button", name="Import & process", exact=True).click()
            expect(page.get_by_role("heading", name="Processing report", exact=True)).to_be_visible(
                timeout=20000
            )
            expect(page.get_by_text("Saving this report in browser history…")).to_have_count(0)

        def download(label):
            with page.expect_download() as event:
                page.get_by_role("button", name=label, exact=True).click()
            path = Path(temporary) / event.value.suggested_filename
            event.value.save_as(path)
            assert path.stat().st_size > 0
            return path

        navigate("/", "Good data. Clear decisions.")
        navigate("/demo/issues", "Your sales overview")
        expect(page.get_by_text("3,445.00", exact=True)).to_be_visible()
        page.get_by_role("tab", name="Issues").click()
        expect(page.get_by_text("normalized_duplicate", exact=True)).to_be_visible()
        navigate("/history", "Processing history")
        expect(page.get_by_text("No saved imports yet", exact=True)).to_be_visible()
        source = (ROOT / "demo/public/demo/issues-source.csv").read_bytes()
        select_file(source)
        process()
        report_route = page.url.split("#")[1]
        expect(page.get_by_text("3,445.00", exact=True)).to_be_visible()
        page.get_by_role("tab", name="Issues").click()
        expect(page.get_by_text("normalized_duplicate", exact=True)).to_be_visible()
        page.get_by_role("tab", name="Issues").focus()
        page.keyboard.press("ArrowRight")
        expect(page.get_by_role("tab", name="Changes")).to_be_focused()
        expect(page.get_by_text("Trimmed surrounding whitespace.")).to_be_visible()
        page.get_by_role("tab", name="Cleaned data").click()
        page.get_by_label("Search cleaned data").fill("no-such-customer")
        expect(page.get_by_text("No matching order lines", exact=True)).to_be_visible()
        page.get_by_label("Search cleaned data").fill("")
        page.get_by_role("button", name="Next report page").click()
        expect(page.get_by_text("2 / 5", exact=True)).to_be_visible()
        cleaned_csv = download("CSV")
        cleaned_xlsx = download("XLSX")
        issues_csv = download("Issues CSV")
        assert len(list(csv.DictReader(StringIO(cleaned_csv.read_text("utf-8-sig"))))) == 48
        assert len(list(csv.DictReader(StringIO(issues_csv.read_text("utf-8-sig"))))) == 8
        assert "USD" in cleaned_xlsx.name
        workbook = load_workbook(cleaned_xlsx, data_only=False)
        assert workbook.active.max_row == 49
        assert workbook.active["E2"].data_type == "n"
        assert workbook.active["F2"].data_type == "n"
        workbook.close()
        select_file(cleaned_xlsx.read_bytes(), "roundtrip.xlsx", "EUR")
        process()
        expect(page.get_by_text("3,445.00", exact=True)).to_be_visible()
        expect(page.get_by_role("tab", name="Issues")).to_contain_text("0")
        assert "EUR" in download("CSV").name
        page.reload()
        expect(page.get_by_role("heading", name="Processing report")).to_be_visible()
        navigate("/history", "Processing history")
        page.get_by_role("link", name="my-orders.csv").click()
        expect(page.get_by_role("heading", name="Processing report")).to_be_visible()
        assert page.url.endswith(report_route)
        results["checks"].extend(
            [
                "examples stay out of history",
                "built worker import",
                "tabs and keyboard",
                "search and pagination",
                "CSV and issues exports",
                "CSV → report → XLSX → reimport",
                "reload and reopen history",
            ]
        )

        header = b"order_id,date,customer,product,quantity,unit_price\n"
        select_file(header + b"1,bad,A,B,0,2", "invalid.csv", "PLN")
        process()
        expect(page.get_by_text("No rows can be exported yet.")).to_be_visible()
        expect(page.get_by_role("button", name="CSV", exact=True)).to_be_disabled()
        download("Issues CSV")
        select_file(b"bad,columns\n1,2", "bad.csv")
        page.get_by_role("button", name="Import & process").click()
        expect(page.get_by_role("alert")).to_contain_text("Missing columns")
        select_file(b"x" * (5 * 1024 * 1024 + 1), "large.csv")
        expect(page.get_by_role("alert")).to_contain_text("too large")
        select_file(header + b"1,2026-01-01,A,B,1,2\n" * 10001, "rows.csv")
        page.get_by_role("button", name="Import & process").click()
        expect(page.get_by_role("alert")).to_contain_text("10,000")
        results["checks"].append("zero valid rows, malformed input and limits")

        held = []
        page.route(re.compile(r"processor\.worker.*\.js"), lambda route: held.append(route))
        select_file(source, "cancelled.csv")
        page.get_by_role("button", name="Import & process").click()
        page.get_by_role("button", name="Cancel processing").click()
        expect(page.get_by_role("alert")).to_contain_text("Processing cancelled")
        for route in held:
            route.abort()
        page.unroute(re.compile(r"processor\.worker.*\.js"))
        navigate("/history", "Processing history")
        expect(page.get_by_role("link", name="cancelled.csv")).to_have_count(0)
        results["checks"].append("cancellation leaves no history entry")

        # Fill the cap through real imports, then check persistence and deletion.
        for index in range(10):
            select_file(
                header + f"{index},2026-01-01,A,B,1,2\n".encode(),
                f"recent-{index}.csv",
                "GBP",
            )
            process()
        navigate("/history", "Processing history")
        expect(page.locator("tbody tr")).to_have_count(10)
        expect(page.get_by_role("link", name="my-orders.csv")).to_have_count(0)
        page.reload()
        expect(page.locator("tbody tr")).to_have_count(10)
        page.on("dialog", lambda dialog: dialog.accept())
        page.get_by_role("button", name="Delete recent-9.csv", exact=True).click()
        expect(page.locator("tbody tr")).to_have_count(9)
        results["checks"].append("latest ten imports and individual deletion")

        axe = os.environ.get("AXE_PATH")
        for width in [320, 375, 768, 1440]:
            page.set_viewport_size({"width": width, "height": 900})
            for route, heading in [
                ("/", "Good data. Clear decisions."),
                ("/import", "Import your sales data"),
                ("/history", "Processing history"),
                ("/demo/issues", "Your sales overview"),
            ]:
                navigate(route, heading)
                if route.startswith("/demo"):
                    expect(page.get_by_role("heading", name="Processing report")).to_be_visible()
                assert page.evaluate("document.documentElement.scrollWidth <= innerWidth"), (
                    width,
                    route,
                )
                if axe:
                    page.add_script_tag(path=axe)
                    violations = page.evaluate(
                        "async () => (await axe.run()).violations.map(v => "
                        "({id:v.id, nodes:v.nodes.map(n=>n.target)}))"
                    )
                    assert not violations, (width, route, violations)
                    results["accessibility"].append({"width": width, "route": route})
                if route == "/demo/issues":
                    for tab in ["Cleaned data", "Issues", "Changes"]:
                        page.get_by_role("tab", name=tab).click()
                        assert page.evaluate(
                            "document.documentElement.scrollWidth <= innerWidth"
                        ), (width, tab)
            results["viewports"].append(width)
            page.locator("#main-content").focus()
            page.screenshot(path=str(OUTPUT / f"report-{width}.png"), full_page=True)
        page.set_viewport_size({"width": 375, "height": 812})
        page.get_by_role("button", name="Open navigation").click()
        expect(page.get_by_role("link", name="FlowOps home")).to_be_focused()
        page.keyboard.press("Escape")
        expect(page.get_by_role("button", name="Open navigation")).to_be_focused()
        page.get_by_role("button", name="Open navigation").click()
        page.get_by_role("navigation").get_by_role("link", name="Import data", exact=True).click()
        expect(page.get_by_role("heading", name="Import your sales data")).to_be_visible()
        expect(page.get_by_role("button", name="Open navigation")).to_have_attribute(
            "aria-expanded", "false"
        )
        results["checks"].append("mobile navigation and focus")
        navigate("/history", "Processing history")
        page.get_by_role("button", name="Clear history", exact=True).click()
        expect(page.get_by_text("No saved imports yet", exact=True)).to_be_visible()
        page.reload()
        expect(page.get_by_text("No saved imports yet", exact=True)).to_be_visible()
        results["checks"].append("clear history survives reload")

        # A storage failure must not discard the result or prevent export.
        context.add_init_script(
            "Object.defineProperty(window, 'indexedDB', "
            "{get(){throw new DOMException('Storage disabled', 'SecurityError')}})"
        )
        page.reload()
        select_file(source, "unsaved.csv")
        process()
        expect(
            page.get_by_text(re.compile("This browser could not save your history"))
        ).to_be_visible()
        assert "USD" in download("XLSX").name
        results["checks"].append("storage warning keeps report and exports usable")
        assert all(method == "GET" and url.startswith(BASE + "/") for method, url in requests), (
            requests
        )
        assert results["workers"] and all(
            "/assets/processor.worker-" in url for url in results["workers"]
        )
        assert not results["errors"], results["errors"]
        results["checks"].append("no file upload or external data requests")
        browser.close()
    (OUTPUT / "results.json").write_text(json.dumps(results, indent=2) + "\n", encoding="utf-8")
    print(
        json.dumps(
            {
                "checks": results["checks"],
                "viewports": results["viewports"],
                "axe_scans": len(results["accessibility"]),
                "worker_runs": len(results["workers"]),
                "errors": results["errors"],
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    run()
