import csv
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import PurePosixPath
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

import pandas as pd
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.config import settings
from app.errors import ImportErrorDetail

COLUMNS = ["order_id", "date", "customer", "product", "quantity", "unit_price"]


@dataclass
class Sheet:
    frame: pd.DataFrame
    source_rows: list[int]
    formulas: dict[int, list[str]]
    column_changes: list[str]


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def read_sheet(content: bytes, filename: str) -> Sheet:
    suffix = PurePosixPath(filename.lower()).suffix
    if suffix not in {".csv", ".xlsx"}:
        raise ImportErrorDetail("unsupported_file", "Choose a CSV or XLSX file.", 415)
    if len(content) > settings.max_file_bytes:
        raise ImportErrorDetail("file_too_large", "Files must be 5 MB or smaller.", 413)
    if not content:
        raise ImportErrorDetail("empty_file", "The uploaded file is empty.")
    rows, source_rows, formula_cells = read_csv(content) if suffix == ".csv" else read_xlsx(content)
    if not rows:
        raise ImportErrorDetail("empty_file", "The file does not contain a header row.")
    original = rows.pop(0)
    source_rows.pop(0)
    headers = [value.strip().lower() for value in original]
    missing = sorted(set(COLUMNS) - set(headers))
    extra = sorted(set(headers) - set(COLUMNS))
    if missing or extra or len(headers) != len(set(headers)):
        details = []
        if missing:
            details.append(f"Missing columns: {', '.join(missing)}.")
        if extra:
            details.append("Unexpected columns. Use only the six required columns.")
        if len(headers) != len(set(headers)):
            details.append("Column names must be unique after trimming and lowercasing.")
        raise ImportErrorDetail("invalid_columns", " ".join(details))
    if not rows:
        raise ImportErrorDetail("no_records", "Add at least one order line below the header.")
    if len(rows) > settings.max_rows:
        raise ImportErrorDetail("too_many_rows", "Files can contain at most 10,000 order lines.")
    for row, number in zip(rows, source_rows, strict=True):
        if not row:
            row.extend([""] * len(headers))
        if len(row) != len(headers):
            raise ImportErrorDetail(
                "malformed_row", f"Row {number} has {len(row)} cells; expected six."
            )
    formulas = {
        row: [headers[column] for column in columns if column < len(headers)]
        for row, columns in formula_cells.items()
    }
    column_changes = [
        f"Column {index + 1} normalized to {normalized}."
        for index, (raw, normalized) in enumerate(zip(original, headers, strict=True))
        if raw != normalized
    ]
    return Sheet(pd.DataFrame(rows, columns=headers), source_rows, formulas, column_changes)


def read_csv(content: bytes) -> tuple[list[list[str]], list[int], dict[int, list[int]]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ImportErrorDetail("invalid_encoding", "Save the CSV as UTF-8 and try again.") from exc
    if "\x00" in text:
        raise ImportErrorDetail("malformed_csv", "The CSV contains invalid null characters.")
    reader = csv.reader(StringIO(text, newline=""), strict=True)
    rows, numbers = [], []
    try:
        while True:
            start_line = reader.line_num + 1
            row = next(reader, None)
            if row is None:
                break
            rows.append(row)
            numbers.append(start_line)
            if len(rows) > settings.max_rows + 1:
                raise ImportErrorDetail(
                    "too_many_rows", "Files can contain at most 10,000 order lines."
                )
    except csv.Error as exc:
        raise ImportErrorDetail(
            "malformed_csv", f"Malformed CSV near row {reader.line_num}."
        ) from exc
    return rows, numbers, {}


def read_xlsx(content: bytes) -> tuple[list[list[str]], list[int], dict[int, list[int]]]:
    try:
        with ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            if (
                len(members) > 200
                or sum(m.file_size for m in members) > settings.max_expanded_bytes
            ):
                raise ImportErrorDetail(
                    "expanded_file_too_large", "The expanded workbook is too large.", 413
                )
            if any(m.flag_bits & 1 for m in members):
                raise ImportErrorDetail(
                    "encrypted_workbook", "Password-protected workbooks are not supported."
                )
        workbook = load_workbook(
            BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        try:
            if len(workbook.worksheets) != 1:
                raise ImportErrorDetail(
                    "multiple_sheets", "Use a workbook with exactly one worksheet."
                )
            sheet = workbook.worksheets[0]
            # Ignore untrusted dimension metadata; bound iteration ourselves.
            sheet.reset_dimensions()
            rows: list[list[str]] = []
            formulas: dict[int, list[int]] = {}
            for number, cells in enumerate(sheet.iter_rows(), start=1):
                if number > settings.max_rows + 1:
                    raise ImportErrorDetail(
                        "too_many_rows", "Files can contain at most 10,000 order lines."
                    )
                if len(cells) > len(COLUMNS):
                    raise ImportErrorDetail("invalid_columns", "Use only the six required columns.")
                row = [cell_text(cell.value) for cell in cells]
                row.extend([""] * (len(COLUMNS) - len(row)))
                rows.append(row)
                formula_columns = [i for i, cell in enumerate(cells) if cell.data_type == "f"]
                if formula_columns:
                    formulas[number] = formula_columns
            return rows, list(range(1, len(rows) + 1)), formulas
        finally:
            workbook.close()
    except ImportErrorDetail:
        raise
    except (
        BadZipFile,
        InvalidFileException,
        KeyError,
        IndexError,
        NotImplementedError,
        ValueError,
        TypeError,
        OSError,
        DefusedXmlException,
        ParseError,
        OverflowError,
    ) as exc:
        raise ImportErrorDetail(
            "malformed_xlsx", "This file is not a readable XLSX workbook."
        ) from exc
