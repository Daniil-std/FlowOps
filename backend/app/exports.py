import csv
import re
from decimal import Decimal
from io import BytesIO, StringIO
from typing import Any

from openpyxl import Workbook

from app.reading import COLUMNS


def export_issues(issues: list[dict[str, Any]]) -> bytes:
    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(["source_row", "column", "code", "explanation"])
    for issue in issues:
        writer.writerow([csv_text(issue[key]) for key in ["row", "column", "code", "message"]])
    return output.getvalue().encode("utf-8-sig")


def safe_filename(filename: str) -> str:
    basename = filename.replace("\\", "/").rsplit("/", 1)[-1]
    return re.sub(r"[^A-Za-z0-9._ -]", "_", basename).strip(" .")[:180] or "dataset.csv"


def csv_text(value: object) -> object:
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def export_csv(records: list[dict[str, Any]]) -> bytes:
    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(COLUMNS)
    for record in records:
        writer.writerow([csv_text(record[column]) for column in COLUMNS])
    return output.getvalue().encode("utf-8-sig")


def export_xlsx(records: list[dict[str, Any]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Cleaned orders"
    sheet.append(COLUMNS)
    for row_number, record in enumerate(records, start=2):
        for column_number, column in enumerate(COLUMNS, start=1):
            value = record[column]
            if column == "unit_price":
                cell = sheet.cell(row_number, column_number, Decimal(value))
                cell.number_format = "0.00"
            elif column == "quantity":
                sheet.cell(row_number, column_number, value)
            else:
                cell = sheet.cell(row_number, column_number, str(value))
                # Explicit text cells preserve identifiers and prevent formula execution.
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in zip("ABCDEF", [20, 16, 26, 30, 14, 16], strict=True):
        sheet.column_dimensions[column].width = width
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
