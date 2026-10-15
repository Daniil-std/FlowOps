import csv
from datetime import datetime
from io import BytesIO, StringIO
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from app.analytics import calculate_analytics
from app.config import settings
from app.errors import ImportErrorDetail
from app.exports import export_csv, export_xlsx, safe_filename
from app.processing import process_file
from app.reading import COLUMNS
from app.sample import sample_csv

VALID = ["ORD-001", "2026-09-01", "Mira Ellis", "Notebook", "2", "18.50"]


def csv_file(rows: list[list[object]], headers: list[str] = COLUMNS) -> bytes:
    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(headers)
    writer.writerows(rows)
    return output.getvalue().encode("utf-8")


def xlsx_file(rows: list[list[object]], headers: list[str] = COLUMNS) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_valid_csv_preserves_ids_and_literal_na() -> None:
    result = process_file(
        csv_file([["001", "2026-09-01", "NA", "NULL", "2", "18.50"]]), "orders.csv"
    )
    assert result.total_rows == 1
    assert result.status == "ready"
    assert result.records[0].order_id == "001"
    assert result.records[0].customer == "NA"
    assert result.records[0].product == "NULL"


def test_valid_xlsx_with_excel_date() -> None:
    row = list(VALID)
    content = xlsx_file([[row[0], datetime(2026, 9, 1), *row[2:]]])
    result = process_file(content, "orders.XLSX")
    assert len(result.records) == 1
    assert str(result.records[0].date) == "2026-09-01"
    assert result.records[0].source_row == 2


@pytest.mark.parametrize(
    "headers", [COLUMNS[:-1], [*COLUMNS, "extra"], [*COLUMNS[:-1], " QUANTITY "]]
)
def test_invalid_columns(headers: list[str]) -> None:
    with pytest.raises(ImportErrorDetail, match="columns|Column"):
        process_file(csv_file([VALID], headers), "orders.csv")


@pytest.mark.parametrize(
    "column,value",
    [
        ("date", "01/09/2026"),
        ("date", "2026-02-30"),
        ("date", "2026-01/01"),
        ("date", "2026-01-01T12:30:00"),
        ("quantity", "0"),
        ("quantity", "-2"),
        ("quantity", "1.5"),
        ("quantity", "ten"),
        ("quantity", "1000001"),
        ("unit_price", "-0.01"),
        ("unit_price", "NaN"),
        ("unit_price", "inf"),
        ("unit_price", "1,23"),
        ("unit_price", "1,234"),
        ("unit_price", "1.234"),
        ("unit_price", "1000001"),
        ("unit_price", "1e2"),
        ("product", "x" * 201),
        ("customer", "a\x01b"),
    ],
)
def test_invalid_values(column: str, value: str) -> None:
    row = list(VALID)
    row[COLUMNS.index(column)] = value
    result = process_file(csv_file([row]), "orders.csv")
    assert result.invalid_rows == 1
    assert result.records == []
    assert result.issues[0].row == 2
    assert result.issues[0].column == column


@pytest.mark.parametrize("column", COLUMNS)
def test_missing_required_values(column: str) -> None:
    row = list(VALID)
    row[COLUMNS.index(column)] = "  "
    result = process_file(csv_file([row]), "orders.csv")
    assert result.invalid_rows == 1
    assert result.issues[0].message == "A required value is missing."


def test_duplicate_lines_and_distinct_items_on_same_order() -> None:
    other_product = [*VALID[:3], "Planner", "1", "28.00"]
    result = process_file(csv_file([VALID, VALID, other_product]), "orders.csv")
    assert (result.total_rows, len(result.records), result.duplicate_rows, result.invalid_rows) == (
        3,
        2,
        1,
        0,
    )
    assert result.issues[0].row == 3
    assert "row 2" in result.issues[0].message
    analytics = calculate_analytics([row.model_dump(mode="json") for row in result.records])
    assert analytics.revenue == "65.00"
    assert analytics.orders == 1
    assert analytics.average_order_value == "65.00"
    assert analytics.units == 3
    assert analytics.top_products[0].product == "Notebook"


def test_normalization_report_and_normalized_duplicate_policy() -> None:
    rows = [
        [" 001 ", "2026/09/01", " Mira ", " Notebook ", "+2.0", "1,234.50"],
        ["001", "2026-09-01", "Mira", "Notebook", "2", "1234.50"],
    ]
    result = process_file(csv_file(rows, [f" {name.upper()} " for name in COLUMNS]), "orders.csv")
    assert len(result.column_changes) == 6
    assert len(result.records) == 1
    assert result.normalized_rows == 1
    assert result.duplicate_rows == 1
    assert result.issues[0].code == "normalized_duplicate"
    assert result.records[0].unit_price == "1234.50"
    assert result.records[0].quantity == 2
    assert {change.column for change in result.changes} == set(COLUMNS)


def test_blank_rows_and_multiline_csv_retain_source_line_numbers() -> None:
    content = (
        b"order_id,date,customer,product,quantity,unit_price\n"
        b'1,2026-09-01,"Mira\nEllis",Book,1,2\n\n2,bad,Mira,Book,1,2\n'
    )
    result = process_file(content, "orders.csv")
    assert result.total_rows == 3
    assert result.records[0].source_row == 2
    assert {issue.row for issue in result.issues} == {4, 5}


def test_zero_price_and_exact_rounding() -> None:
    result = process_file(
        csv_file(
            [
                ["1", *VALID[1:4], "3", "0.10"],
                ["2", *VALID[1:4], "1", "0.00"],
                ["3", *VALID[1:4], "1", "0.01"],
            ]
        ),
        "orders.csv",
    )
    analytics = calculate_analytics([row.model_dump(mode="json") for row in result.records])
    assert analytics.revenue == "0.31"
    assert analytics.orders == 3
    assert analytics.average_order_value == "0.10"
    assert analytics.revenue_over_time[0].revenue == "0.31"
    assert calculate_analytics([]).orders == 0


def test_half_cent_average_rounds_up() -> None:
    result = process_file(
        csv_file([["1", *VALID[1:4], "1", "0.01"], ["2", *VALID[1:4], "1", "0.00"]]), "orders.csv"
    )
    assert (
        calculate_analytics(
            [row.model_dump(mode="json") for row in result.records]
        ).average_order_value
        == "0.01"
    )


def test_xlsx_formulas_are_excluded() -> None:
    result = process_file(xlsx_file([[*VALID[:5], "=1+1"]]), "orders.xlsx")
    assert result.invalid_rows == 1
    assert "formulas" in result.issues[0].message


def test_exports_contain_processed_rows_and_neutralize_formulas() -> None:
    result = process_file(
        csv_file(
            [["001", "2026-09-01", "=SUM(1)", "@danger", "2", "1.25"], ["2", *VALID[1:4], "0", "2"]]
        ),
        "orders.csv",
    )
    records = [row.model_dump(mode="json") for row in result.records]
    exported_csv = list(csv.DictReader(StringIO(export_csv(records).decode("utf-8-sig"))))
    assert len(exported_csv) == 1
    assert exported_csv[0]["customer"] == "'=SUM(1)"
    assert exported_csv[0]["product"] == "'@danger"
    assert exported_csv[0]["order_id"] == "001"
    workbook = load_workbook(BytesIO(export_xlsx(records)), data_only=False)
    sheet = workbook.active
    assert sheet is not None
    assert sheet.max_row == 2
    assert sheet["C2"].value == "=SUM(1)"
    assert sheet["C2"].data_type == "s"
    assert sheet["A2"].value == "001"
    assert sheet["E2"].value == 2
    assert sheet["F2"].value == 1.25
    workbook.close()


@pytest.mark.parametrize(
    "name,content,code",
    [
        ("bad.txt", b"hello", "unsupported_file"),
        ("bad.csv", b"", "empty_file"),
        ("bad.xlsx", b"not excel", "malformed_xlsx"),
        ("bad.csv", b"\xff\xfe", "invalid_encoding"),
        (
            "bad.csv",
            b'order_id,date,customer,product,quantity,unit_price\n"unterminated',
            "malformed_csv",
        ),
        (
            "bad.csv",
            b"order_id,date,customer,product,quantity,unit_price\n1,2,3,4,5,6,7",
            "malformed_row",
        ),
        ("bad.csv", (",".join(COLUMNS) + "\n").encode(), "no_records"),
    ],
)
def test_malformed_files(name: str, content: bytes, code: str) -> None:
    with pytest.raises(ImportErrorDetail) as error:
        process_file(content, name)
    assert error.value.code == code


def test_oversized_files_and_record_limit(monkeypatch: pytest.MonkeyPatch) -> None:
    with pytest.raises(ImportErrorDetail) as error:
        process_file(b"a" * (settings.max_file_bytes + 1), "orders.csv")
    assert error.value.status == 413
    monkeypatch.setattr(settings, "max_rows", 2)
    for name, content in [
        ("orders.csv", csv_file([VALID] * 3)),
        ("orders.xlsx", xlsx_file([VALID] * 3)),
    ]:
        with pytest.raises(ImportErrorDetail) as error:
            process_file(content, name)
        assert error.value.code == "too_many_rows"


def test_rejects_expanded_archive_before_parsing(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "max_expanded_bytes", 100)
    with pytest.raises(ImportErrorDetail) as error:
        process_file(xlsx_file([VALID]), "orders.xlsx")
    assert error.value.code == "expanded_file_too_large"


def test_rejects_multiple_worksheets() -> None:
    workbook = Workbook()
    workbook.create_sheet("Extra")
    content = BytesIO()
    workbook.save(content)
    with pytest.raises(ImportErrorDetail) as error:
        process_file(content.getvalue(), "orders.xlsx")
    assert error.value.code == "multiple_sheets"


def test_safe_filename_and_deterministic_sample() -> None:
    assert safe_filename('..\\folder/evil\r\n".csv') == "evil___.csv"
    result = process_file(sample_csv(), "example.csv")
    assert len(result.records) == 48
    assert result.issues == []
    analytics = calculate_analytics([row.model_dump(mode="json") for row in result.records])
    assert analytics.orders == 24
    assert len(analytics.top_products) == 5
    assert len(analytics.revenue_over_time) == 24


def test_invalid_zip_content() -> None:
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("hello.txt", "hello")
    with pytest.raises(ImportErrorDetail, match="readable XLSX"):
        process_file(buffer.getvalue(), "orders.xlsx")


@pytest.mark.parametrize("column,value", [(1, "2026-09-02"), (2, "Other customer")])
def test_conflicting_order_identity_excludes_every_item(column: int, value: str) -> None:
    other = list(VALID)
    other[column] = value
    result = process_file(csv_file([VALID, other, VALID]), "orders.csv")
    assert result.records == []
    assert result.invalid_rows == 3
    assert result.duplicate_rows == 0
    assert [(issue.row, issue.code) for issue in result.issues] == [
        (2, "order_conflict"),
        (3, "order_conflict"),
        (4, "order_conflict"),
    ]


def test_conflict_on_otherwise_invalid_item_is_not_silently_ignored() -> None:
    other = [VALID[0], "2026-09-02", *VALID[2:4], "0", VALID[5]]
    result = process_file(csv_file([VALID, other]), "orders.csv")
    assert result.invalid_rows == 2
    assert result.records == []
    assert len(result.issues) == 3


def test_signed_zero_is_the_same_price_and_different_quantities_are_retained() -> None:
    result = process_file(
        csv_file(
            [
                [*VALID[:5], "-0.00"],
                [*VALID[:5], "0"],
                [*VALID[:4], "3", "0.00"],
            ]
        ),
        "orders.csv",
    )
    assert len(result.records) == 2
    assert result.duplicate_rows == 1
    assert result.records[0].unit_price == "0.00"


def test_normalized_duplicates_do_not_inflate_revenue() -> None:
    result = process_file(
        csv_file(
            [
                VALID,
                [" ORD-001 ", "2026/09/01", " Mira Ellis ", "Notebook", "+2.00", "18.5"],
            ]
        ),
        "orders.csv",
    )
    analytics = calculate_analytics([row.model_dump(mode="json") for row in result.records])
    assert analytics.revenue == "37.00"
    assert result.issues[0].code == "normalized_duplicate"
    assert "combine their quantities" in result.issues[0].message


def test_malformed_workbook_xml() -> None:
    output = BytesIO()
    with ZipFile(BytesIO(xlsx_file([VALID]))) as source, ZipFile(output, "w") as target:
        for member in source.infolist():
            content = source.read(member.filename)
            if member.filename == "xl/worksheets/sheet1.xml":
                content = b"<worksheet><broken"
            target.writestr(member, content)
    with pytest.raises(ImportErrorDetail) as error:
        process_file(output.getvalue(), "orders.xlsx")
    assert error.value.code == "malformed_xlsx"


def test_ten_thousand_rows_at_numeric_limits_remain_exact() -> None:
    records = [
        {
            "order_id": str(index),
            "date": "2026-09-01",
            "customer": "Example",
            "product": "Example",
            "quantity": 1_000_000,
            "unit_price": "1000000.00",
        }
        for index in range(10_000)
    ]
    analytics = calculate_analytics(records)
    assert analytics.revenue == "10000000000000000.00"
    assert analytics.average_order_value == "1000000000000.00"
    assert analytics.orders == 10_000
